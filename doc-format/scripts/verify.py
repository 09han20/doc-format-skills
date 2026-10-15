# -*- coding: utf-8 -*-
"""对生成的 .docx 做结构断言 + 内容无损校验。

    python verify.py out.docx              # 只做结构断言
    python verify.py out.docx --md src.md  # 追加内容无损校验

结构断言覆盖那些「肉眼看不出来、但一打开 Word 就发现不对」的项：
真标题样式、eastAsia 字体槽、szCs、firstLineChars、三线表边框、
以及最容易翻车的一条 —— 表格单元格有没有从 Normal 继承到首行缩进。
"""
import argparse
import os
import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# 同 md2docx：中文 Windows 控制台是 GBK，✅/❌ 会让汇总行抛 UnicodeEncodeError
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass

import md_parse as P
import spec

# 生成器给列表项加的前缀。剥的时候注意 (?!\d)：否则 "1.1" 会被吃掉成 "1"
LIST_MARK = re.compile(r'^(?:[·□]|\d+\.(?!\d))')
CAPNUM = re.compile(r'^(表|图|代码|算法)\s*\d+\s*')
TOC_HINT = '更新目录'


class Report:
    def __init__(self):
        self.fail = 0
        self.total = 0

    def check(self, cond, msg):
        self.total += 1
        if not cond:
            self.fail += 1
        print(('  [OK] ' if cond else '  [!!] ') + msg)
        return cond

    def summary(self):
        print('\n%s  %d/%d 项通过' % ('全部通过 ✅' if not self.fail else '存在问题 ❌',
                                     self.total - self.fail, self.total))
        return self.fail


def _v(el, tag, attr='val'):
    if el is None:
        return None
    e = el.find(qn('w:' + tag))
    return e.get(qn('w:' + attr)) if e is not None else None


def check_styles(d, R):
    print('\n=== 样式层 ===')
    for name, ea, sz, lvl in (('Normal', spec.FONT_BODY_CJK, spec.hp(spec.BODY_PT), None),
                              ('Heading 1', spec.FONT_HEAD_CJK, spec.hp(spec.HEAD_PT[1]), '0'),
                              ('Heading 2', spec.FONT_HEAD_CJK, spec.hp(spec.HEAD_PT[2]), '1'),
                              ('Heading 3', spec.FONT_HEAD_CJK, spec.hp(spec.HEAD_PT[3]), '2'),
                              ('表题', spec.FONT_CAP_CJK, spec.hp(spec.CAP_PT), None),
                              ('表格文字', spec.FONT_BODY_CJK, spec.hp(spec.TABLE_PT), None)):
        try:
            el = d.styles[name].element
        except KeyError:
            R.check(False, '%s 样式缺失' % name)
            continue
        rpr, ppr = el.find(qn('w:rPr')), el.find(qn('w:pPr'))
        rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
        got_ea = rf.get(qn('w:eastAsia')) if rf is not None else None
        R.check(got_ea == ea and _v(rpr, 'sz') == str(sz)
                and _v(rpr, 'szCs') == str(sz)
                and (lvl is None or _v(ppr, 'outlineLvl') == lvl),
                '%-9s eastAsia=%s sz=%s szCs=%s outlineLvl=%s'
                % (name, got_ea, _v(rpr, 'sz'), _v(rpr, 'szCs'), _v(ppr, 'outlineLvl')))

    R.check(_v(d.styles['Normal'].element.find(qn('w:pPr')), 'ind', 'firstLineChars') == '200',
            'Normal firstLineChars=200（首行缩进 2 字符）')
    R.check(_v(d.styles['Normal'].element.find(qn('w:pPr')), 'ind', 'firstLine') ==
            str(spec.first_line_twips()), 'Normal firstLine=%d（双写降级值）'
            % spec.first_line_twips())
    R.check(_v(d.styles['Normal'].element.find(qn('w:pPr')), 'spacing', 'line') == '360',
            'Normal line=360（1.5 倍行距）')
    R.check(_v(d.styles['表格文字'].element.find(qn('w:pPr')), 'ind', 'firstLineChars') == '0',
            '表格文字 firstLineChars=0（防缩进漏进表格）')

    # --- 封面标题：必须是自建样式，且不得带段落边框 ---
    try:
        R.check(d.styles[spec.STYLE_TITLE].element.find(qn('w:pPr')) is not None,
                '%s 样式存在' % spec.STYLE_TITLE)
    except KeyError:
        R.check(False, '%s 样式缺失' % spec.STYLE_TITLE)

    # 内置 Title / Caption 自带 accent1 边框（蓝线），必须已剥离且未被引用
    for nm in (spec.STYLE_TITLE, 'Title', 'Caption'):
        try:
            ppr = d.styles[nm].element.find(qn('w:pPr'))
        except KeyError:
            continue
        bdr = ppr.find(qn('w:pBdr')) if ppr is not None else None
        R.check(bdr is None, '%-9s 无段落边框（pBdr），不会画出蓝色下划线' % nm)

    used = sorted({p.style.name for p in d.paragraphs
                   if p.style.name in ('Title', 'Caption')})
    R.check(not used, '正文未引用内置 Title/Caption 样式%s'
            % ('（实际用了 %s）' % '/'.join(used) if used else ''))


def check_tables(d, R):
    print('\n=== 表格 ===')
    bad_border, bad_cell, ncell, styles = [], [], 0, set()
    for ti, t in enumerate(d.tables):
        styles.add(t.style.name)
        tb = t._tbl.tblPr.find(qn('w:tblBorders'))
        vals = {c.tag.split('}')[1]: (c.get(qn('w:val')), c.get(qn('w:sz')))
                for c in tb} if tb is not None else {}
        if (vals.get('top') != ('single', str(spec.RULE_THICK))
                or vals.get('bottom') != ('single', str(spec.RULE_THICK))
                or any(vals.get(k) != ('none', '0')
                       for k in ('left', 'right', 'insideH', 'insideV'))):
            bad_border.append(ti)
        for r in t.rows:
            for c in r.cells:
                ncell += 1
                for p in c.paragraphs:
                    ppr = p._p.find(qn('w:pPr'))
                    if _v(ppr, 'ind', 'firstLineChars') != '0':
                        bad_cell.append((ti, _v(ppr, 'ind', 'firstLineChars')))
    R.check(not bad_border, '全部 %d 张表为三线表（顶/底 %s=%d，其余 none）'
            % (len(d.tables), 'sz', spec.RULE_THICK))
    R.check(styles <= {'Normal Table'}, '表格样式 %s（已脱离 Table Grid 全网格线）' % styles)
    R.check(not bad_cell, '单元格首行缩进全为 0（查了 %d 格，异常 %d）' % (ncell, len(bad_cell)))
    if d.tables:
        t0 = d.tables[0]
        tcpr = t0.rows[0].cells[0]._tc.find(qn('w:tcPr'))
        R.check(_v(tcpr, 'tcBorders', 'sz') is None and
                tcpr.find(qn('w:tcBorders')) is not None and
                tcpr.find(qn('w:tcBorders')).find(qn('w:bottom')).get(qn('w:sz')) == str(spec.RULE_THIN),
                '表头行下有 %d/8 pt 细线' % spec.RULE_THIN)
        R.check(t0.rows[0]._tr.find(qn('w:trPr')).find(qn('w:tblHeader')) is not None,
                '表头行设为跨页重复')

        bad_jc = []
        for ti, t in enumerate(d.tables):
            for i, r in enumerate(t.rows):
                for c in r.cells:
                    for p in c.paragraphs:
                        want = spec.TABLE_HEADER_JC if i == 0 else spec.TABLE_CELL_JC
                        if _v(p._p.find(qn('w:pPr')), 'jc') != want:
                            bad_jc.append((ti, i, _v(p._p.find(qn('w:pPr')), 'jc')))
        R.check(not bad_jc, '单元格水平对齐 表头=%s 表体=%s（异常 %d）'
                % (spec.TABLE_HEADER_JC, spec.TABLE_CELL_JC, len(bad_jc)))
        if bad_jc:
            print('      异常样例:', bad_jc[:4])

        va = t0.cell(0, 0)._tc.find(qn('w:tcPr')).find(qn('w:vAlign'))
        R.check(va is not None and va.get(qn('w:val')) == spec.TABLE_CELL_VALIGN,
                '单元格垂直对齐 = %s' % spec.TABLE_CELL_VALIGN)


def check_page(d, path, R):
    print('\n=== 页面 ===')
    s = d.sections[0]
    R.check(abs(s.left_margin.cm - spec.MARGIN_LEFT_CM) < .02
            and abs(s.right_margin.cm - spec.MARGIN_RIGHT_CM) < .02
            and abs(s.top_margin.cm - spec.MARGIN_TOP_CM) < .02
            and abs(s.bottom_margin.cm - spec.MARGIN_BOTTOM_CM) < .02,
            '页边距 上/下/左/右 = %.1f/%.1f/%.1f/%.1f cm'
            % (s.top_margin.cm, s.bottom_margin.cm, s.left_margin.cm, s.right_margin.cm))
    z = zipfile.ZipFile(path)
    footers = [n for n in z.namelist() if re.match(r'word/footer\d+\.xml', n)]
    R.check(bool(footers) and all(b'PAGE' in z.read(f) for f in footers),
            '页脚含 PAGE 域（%s）' % (', '.join(footers) or '无'))
    body = z.read('word/document.xml')
    has_toc = b'TOC' in body
    R.check(True, 'TOC 域：%s' % ('已插入' if has_toc else '未插入（--toc 才插）'))


def docx_stream(doc):
    """产出 (文本, 是否列表项)。列表项靠悬挂缩进判定 —— 不能靠文本前缀，
    因为表格单元格里合法存在「1.战略冲突管理」这种以编号开头的文字。"""
    out = []
    for ch in doc.element.body.iterchildren():
        if ch.tag == qn('w:p'):
            p = Paragraph(ch, doc)
            if p.text.strip():
                fi = p.paragraph_format.first_line_indent
                out.append((p.text, bool(fi is not None and fi < 0)))
        elif ch.tag == qn('w:tbl'):
            for r in Table(ch, doc).rows:
                for c in r.cells:
                    for p in c.paragraphs:
                        if p.text.strip():
                            out.append((p.text, False))
    return out


def md_stream(blocks, counters=None):
    out = []
    for k, v in blocks:
        if k == 'table':
            if v.get('cap'):
                out.append('%s %s %s' % (v['cap']['kind'], v['cap']['num'] or '',
                                         v['cap']['text']))
            out += [P.strip_inline(c) for r in v['rows'] for c in r]
            if v.get('note'):
                out.append(v['note'])
        elif k == 'img':
            # 图片本身不产出文本；题注在图片**下方**渲染成「图 N 说明」一行，
            # 与表格（表题在上方）一样计入文本流比对。
            if v.get('cap'):
                out.append('%s %s %s' % (v['cap']['kind'], v['cap']['num'] or '',
                                         v['cap']['text']))
            if v.get('note'):
                out.append(v['note'])
        elif k == 'code':
            out.extend(v[1].split('\n'))   # 代码块每行渲染成独立段落，逐行对齐
        elif k in ('h', 'ol'):
            out.append(P.strip_inline(v[1]))
        elif k == 'title':
            out.append(v)
        elif k in ('p', 'quote', 'todo', 'ul'):
            out.append(P.strip_inline(v))
        elif k == 'cap':
            out.append(v.get('raw') or ('%s %s %s' % (v['kind'], v['num'] or '', v['text'])))
        elif k == 'note':
            out.append(v)
    return out


def check_images(d, md_path, R):
    """插图张数与文件名核对。图片没有文本，内容无损那条查不到它 ——
    张数对不上（比如路径写错被跳过）必须在这里暴露。"""
    print('\n=== 插图 ===')
    blocks = P.parse_blocks(md_path, is_path=True)
    imgs = [v for k, v in blocks if k == 'img']
    drawn = len(d.inline_shapes)
    R.check(drawn == len(imgs), '图片张数 md %d vs docx %d' % (len(imgs), drawn))
    # 嵌入后 Word 会把图片改名为 media/imageN.png，比文件名没意义 ——
    # 改比内容指纹：确认 md 里引用的每个图片文件都真的进了包。
    import hashlib
    base = os.path.dirname(os.path.abspath(md_path))
    want = {}
    for v in imgs:
        p = v['src'] if os.path.isabs(v['src']) else os.path.join(base, v['src'])
        if not os.path.isfile(p):
            R.check(False, '图片文件缺失：%s' % v['src'])
            return
        want[hashlib.sha1(open(p, 'rb').read()).hexdigest()] = v['src']
    got = {hashlib.sha1(rel.target_part.blob).hexdigest()
           for rel in d.part.rels.values() if rel.reltype.endswith('/image')}
    R.check(set(want) <= got, '图片内容已嵌入（%s）'
            % (', '.join(sorted(want.values())) or '无'))
    capped = [v for v in imgs if v.get('cap')]
    if capped:
        print('      有题注的图 %d 张：%s'
              % (len(capped), ' / '.join(v['cap']['kind'] + ' ' + v['cap']['text']
                                        for v in capped)))


def check_content(d, md_path, R):
    print('\n=== 内容无损 ===')
    nz = lambda s: re.sub(r'\s+', '', s or '')
    # 自动编号是生成器**添加**的信息，源文件里没有 → 比对时两侧都剥掉号，
    # 否则「表 1 说明」永远对不上「表 说明」。
    canon = lambda x: CAPNUM.sub(r'\1', x)
    blocks = P.parse_blocks(md_path, is_path=True)
    A = [canon(nz(x)) for x in md_stream(blocks)]
    A = [x for x in A if x]
    B = []
    for x, is_list in docx_stream(d):
        if TOC_HINT in x:            # TOC 域的缓存占位文本，不是正文
            continue
        x = nz(LIST_MARK.sub('', x) if is_list else x)
        if x:
            B.append(canon(x))
    R.check(A == B, '按 body 顺序的文本流逐单元一致（md %d vs docx %d）' % (len(A), len(B)))
    if A != B:
        i = next((i for i, (x, y) in enumerate(zip(A, B)) if x != y), min(len(A), len(B)))
        print('      #%d\n        md  =%r\n        docx=%r'
              % (i, (A[i] if i < len(A) else '<无>')[:88],
                 (B[i] if i < len(B) else '<无>')[:88]))
    sa, sb = nz(''.join(A)), nz(''.join(B))
    R.check(sa == sb, '全文去除格式后字符完全一致（%d vs %d）—— 内容零丢失' % (len(sa), len(sb)))
    if sa != sb:
        i = next((i for i, (x, y) in enumerate(zip(sa, sb)) if x != y), min(len(sa), len(sb)))
        print('      @%d md…%r\n         dc…%r'
              % (i, sa[max(0, i - 30):i + 30], sb[max(0, i - 30):i + 30]))


def main(argv=None):
    ap = argparse.ArgumentParser(description='校验生成的 docx')
    ap.add_argument('docx')
    ap.add_argument('--md', help='源 Markdown，给了就做内容无损校验')
    a = ap.parse_args(argv)

    d = Document(a.docx)
    R = Report()
    print('校验 %s' % a.docx)
    check_styles(d, R)
    check_tables(d, R)
    check_page(d, a.docx, R)
    if a.md:
        check_images(d, a.md, R)
        check_content(d, a.md, R)
    print('\n段落 %d | 表格 %d | 节 %d' % (len(d.paragraphs), len(d.tables), len(d.sections)))
    return 1 if R.summary() else 0


if __name__ == '__main__':
    sys.exit(main())