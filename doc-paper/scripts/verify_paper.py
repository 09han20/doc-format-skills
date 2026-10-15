# -*- coding: utf-8 -*-
"""对 md2paper 生成的论文 docx 做结构断言 + 内容无损校验。

    python verify_paper.py paper.docx --md 论文.md
    python verify_paper.py daily.docx --md 论文.md --mode daily

模式优先从 docx 核心属性里读（md2paper 生成时写进去的），读不到再用 --mode。

断言里最要紧的两条：
  · 分节与页码 —— 封面无页脚、前置罗马、正文阿拉伯从 1 起
  · 内容无损   —— 源 md 与产物的全文（剥格式后）逐字符一致；改排版不能改掉字
"""
import argparse
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8')
    except Exception:
        pass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from _engine import P, V, spec as S
import latex_omml as L
import md2paper
import md_parse_paper as PP
import spec_paper as SP

# 自动编号是生成器**添加**的信息，源稿里没有 → 比对时两侧都剥掉号
CAPNUM = re.compile(r'^(表|图|公式|式|算法|代码)\s*\d+(?:-\d+)?\s*')
EQNUM = re.compile(r'\(\d{1,2}(?:-\d{1,3})?\)')

EXPECT_STYLES = (
    #        样式名                中文字体            字号pt            outlineLvl 对齐
    ('Normal', S.FONT_BODY_CJK, S.BODY_PT, None, None),
    ('Heading 1', SP.FONT_H[1], SP.HEAD_PT[1], '0', 'center'),
    ('Heading 2', SP.FONT_H[2], SP.HEAD_PT[2], '1', 'left'),
    ('Heading 3', SP.FONT_H[3], SP.HEAD_PT[3], '2', 'left'),
    ('Heading 4', SP.FONT_H[4], SP.HEAD_PT[4], '3', 'left'),
    (SP.STYLE_COVER_TITLE, SP.FONT_COVER_TITLE, SP.COVER_TITLE_PT, None, 'center'),
    (SP.STYLE_TOC_TITLE, S.FONT_HEAD_CJK, SP.TOC_TITLE_PT, None, 'center'),
    (SP.STYLE_CAPTION, SP.FONT_CAP, SP.CAP_PT, None, 'center'),
    (SP.STYLE_TBL, S.FONT_BODY_CJK, S.TABLE_PT, None, None),
    (SP.STYLE_REF, S.FONT_BODY_CJK, SP.REF_PT, None, 'both'),
)
EXPECT_BOLD = ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4',
               SP.STYLE_COVER_TITLE, SP.STYLE_CAPTION, SP.STYLE_TOC_TITLE)


def _search(pat, s):
    """从核心属性串里抠出标记值（模式 / 公式形式都是生成时写进去的）。"""
    m = re.search(pat, s or '')
    return m.group(1) if m else None


def _rpr(d, name):
    el = d.styles[name].element
    return el.find(qn('w:rPr'))


def check_styles(d, R):
    print('\n=== 样式层 ===')
    for name, ea, pt, lvl, jc in EXPECT_STYLES:
        try:
            el = d.styles[name].element
        except KeyError:
            R.check(False, '%s 样式缺失' % name)
            continue
        rpr, ppr = el.find(qn('w:rPr')), el.find(qn('w:pPr'))
        rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
        got_ea = rf.get(qn('w:eastAsia')) if rf is not None else None
        ok = (got_ea == ea and V._v(rpr, 'sz') == str(S.hp(pt))
              and V._v(rpr, 'szCs') == str(S.hp(pt))
              and (lvl is None or V._v(ppr, 'outlineLvl') == lvl)
              and (jc is None or V._v(ppr, 'jc') == jc))
        R.check(ok, '%-7s eastAsia=%s sz=%s szCs=%s outlineLvl=%s jc=%s'
                % (name, got_ea, V._v(rpr, 'sz'), V._v(rpr, 'szCs'),
                   V._v(ppr, 'outlineLvl'), V._v(ppr, 'jc')))
    for name in EXPECT_BOLD:
        R.check(V._v(_rpr(d, name), 'b') == '1', '%-7s 加粗（md 规定）' % name)

    # 正文与不缩进的样式
    R.check(V._v(d.styles['Normal'].element.find(qn('w:pPr')), 'ind', 'firstLineChars')
            == '200', 'Normal firstLineChars=200（首行缩进 2 字符）')
    R.check(V._v(d.styles['Normal'].element.find(qn('w:pPr')), 'spacing', 'line')
            == '360', 'Normal line=360（1.5 倍行距）')
    for nm in (SP.STYLE_TBL, SP.STYLE_KW, SP.STYLE_EQ, SP.STYLE_TOC_TITLE):
        R.check(V._v(d.styles[nm].element.find(qn('w:pPr')), 'ind', 'firstLineChars')
                == '0', '%s firstLineChars=0（不继承正文缩进）' % nm)

    # 参考文献悬挂缩进
    ind = d.styles[SP.STYLE_REF].element.find(qn('w:pPr')).find(qn('w:ind'))
    hanging = int(ind.get(qn('w:hanging')) or 0)
    R.check(hanging > 0 and ind.get(qn('w:firstLineChars')) == '0',
            '参考文献 悬挂缩进 %d twips（首行显式回 0）' % hanging)

    # 封面标题不得带内置 Title 的蓝色下边框
    for nm in (SP.STYLE_COVER_TITLE, SP.STYLE_TITLE, 'Title', 'Caption'):
        try:
            ppr = d.styles[nm].element.find(qn('w:pPr'))
        except KeyError:
            continue
        R.check(ppr is None or ppr.find(qn('w:pBdr')) is None,
                '%-7s 无段落边框（pBdr），不会画出蓝色下划线' % nm)


def check_sections(d, path, blocks, st, R, mode):
    """分节与页码：封面节无页脚 → 前置 upperRoman → 正文 decimal 从 1 起。"""
    print('\n=== 分节与页码 ===')
    zones = PP.plan_zones(blocks, st, mode)
    want = 1 + (1 if st['has_meta'] else 0) + (1 if zones['front'] else 0)
    secs = d.sections
    R.check(len(secs) == want,
            '节数 %d（封面%s 前置%s 正文1）' % (want,
                                            '有' if st['has_meta'] else '无',
                                            '有' if zones['front'] else '无'))
    if len(secs) != want:
        return
    i = 0
    if st['has_meta']:
        sp = secs[i]._sectPr
        R.check(not sp.findall(qn('w:footerReference')),
                '封面节无页脚引用（封面不编页码）')
        i += 1
    if zones['front']:
        sp = secs[i]._sectPr
        pn = sp.find(qn('w:pgNumType'))
        R.check(pn is not None and pn.get(qn('w:fmt')) == SP.PAGE_FMT_FRONT
                and pn.get(qn('w:start')) == str(SP.PAGE_START),
                '前置节 页码格式 %s，从 %s 起' % (SP.PAGE_FMT_FRONT, SP.PAGE_START))
        i += 1
    sp = secs[i]._sectPr
    pn = sp.find(qn('w:pgNumType'))
    R.check(pn is not None and pn.get(qn('w:fmt')) == SP.PAGE_FMT_MAIN
            and pn.get(qn('w:start')) == str(SP.PAGE_START),
            '正文节 页码格式 %s，从 %s 起（阿拉伯数字重新数）'
            % (SP.PAGE_FMT_MAIN, SP.PAGE_START))
    # pgNumType 必须排在 cols 之前 —— 插错位置 Word 会报「文件内容有问题」
    R.check([c.tag.split('}')[1] for c in sp].index('pgNumType')
            < [c.tag.split('}')[1] for c in sp].index('cols'),
            'pgNumType 的 schema 位置正确（在 cols 之前）')

    z = zipfile.ZipFile(path)
    footers = [n for n in z.namelist() if re.match(r'word/footer\d+\.xml', n)]
    R.check(len(footers) == (1 if zones['front'] else 0) + 1,
            '页脚部件 %d 个 = 需要页码的节数（封面那节不建页脚）' % len(footers))
    R.check(all(b'PAGE' in z.read(f) for f in footers),
            '页脚都含 PAGE 域（%s）' % (', '.join(footers) or '无'))


def check_cover(d, st, R):
    print('\n=== 封面 ===')
    if not st['has_meta']:
        R.check(True, '无 front matter，不出封面（日常版允许）')
        return
    styles = [p.style.name for p in d.paragraphs[:12]]
    R.check(SP.STYLE_COVER_TITLE in styles, '封面题名用自建样式（无蓝线）')
    texts = [p.text for p in d.paragraphs[:12]]
    R.check(bool(st['meta_title']) and st['meta_title'] in texts,
            '封面题名已排上：%s' % (st['meta_title'] or '（缺）'))
    missing = [ln for ln in st['cover_entries'] if ln not in texts]
    R.check(not missing, '封面字段行齐全（%d 行）%s'
            % (len(st['cover_entries']), '缺：%s' % missing if missing else ''))


def check_kw(d, R, mode):
    print('\n=== 摘要与关键词 ===')
    kws = [p for p in d.paragraphs if p.style.name == SP.STYLE_KW]
    if mode == SP.MODE_DAILY:
        R.check(not kws, '日常版已省略关键词行（0 行）')
        return
    R.check(bool(kws), '有关键词行（%d 行）' % len(kws))
    bad = []
    for p in kws:
        runs = p._p.findall(qn('w:r'))
        if not runs:
            bad.append(p.text[:20])
            continue
        rpr = runs[0].find(qn('w:rPr'))
        rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
        if V._v(rpr, 'b') != '1' or (rf is None or
                                     rf.get(qn('w:eastAsia')) != SP.FONT_KW_LABEL):
            bad.append(p.text[:20])
    R.check(not bad, '关键词标签为黑体加粗（异常 %d）%s'
            % (len(bad), bad[:3] if bad else ''))


def check_refs(d, st, R, mode):
    print('\n=== 参考文献 ===')
    paras = [p for p in d.paragraphs if p.style.name == SP.STYLE_REF]
    want = 0 if mode == SP.MODE_DAILY else st['n_refs']
    R.check(len(paras) == want,
            '参考文献条目数 %d（应为 %d%s）'
            % (len(paras), want, '，日常版省略' if mode == SP.MODE_DAILY else ''))
    if not paras:
        return
    bad = []
    for p in paras:
        ind = p._p.find(qn('w:pPr')).find(qn('w:ind'))
        if ind is None or int(ind.get(qn('w:hanging')) or 0) <= 0:
            bad.append(p.text[:20])
    R.check(not bad, '每条都有悬挂缩进（异常 %d）' % len(bad))
    nums = [re.match(r'^\[(\d+)\]', p.text) for p in paras]
    R.check(all(nums) and [int(m.group(1)) for m in nums] == list(
        range(1, len(paras) + 1)),
        '编号 [1]…[%d] 连续' % len(paras))


def _math_text(el):
    """把一个元素里所有 m:t 拼起来（数学区的可见字符）。"""
    return ''.join(t.text or '' for t in el.iter(qn('m:t')))


class _TableView:
    """只把 body_tables 暴露成 .tables 的替身 —— doc-format 的 check_tables
    只认 `d.tables`，这样能把公式表摘出去后又原样复用那套三线表断言。"""

    def __init__(self, d, tables):
        self._d = d
        self.tables = tables


def formula_tables(d):
    """认出公式表：1 行 3 列、中栏段落用「公式」样式（生成器就是这么排的）。

    公式不放在正文段落里，是因为 Word 只在 oMathPara 里按**显示式**排公式 ——
    而 oMathPara 与普通 run 同段就失效，只能独占表格中栏。
    """
    out = []
    for t in d.tables:
        try:
            if len(t.rows) != 1 or len(t.columns) != 3:
                continue
            if any(p.style.name == SP.STYLE_EQ for p in t.cell(0, 1).paragraphs):
                out.append(t)
        except Exception:
            continue
    return out


def check_formula(d, R, formula_mode='omml'):
    print('\n=== 公式 ===')
    eqs = formula_tables(d)
    if not eqs:
        R.check(True, '无公式（源稿里没有 $$ … $$）')
        return
    numbered = n_math = 0
    bad, font_bad, align_bad = [], [], []
    for t in eqs:
        body = t.cell(0, 1).paragraphs[0]
        num_cell = t.cell(0, 2)
        num_txt = ''.join(p.text for p in num_cell.paragraphs)
        if num_txt.strip():                  # 空编号栏（$$! 不编号）不参与断言
            numbered += 1
            if (not EQNUM.fullmatch(num_txt.strip())
                    or num_cell.paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.RIGHT):
                align_bad.append(num_txt[:12])
        para = body._p.find(qn('m:oMathPara'))
        math = para.find(qn('m:oMath')) if para is not None else None
        if math is not None:
            n_math += 1
            rf = math.find('.//' + qn('w:rFonts'))
            if rf is None or rf.get(qn('w:ascii')) != L.MATH_FONT:
                font_bad.append(body.text[:12])
    R.check(bool(eqs), '有公式（共 %d 条，1×3 无框表格版式）' % len(eqs))
    R.check(not align_bad, '编号右对齐且形如 (2-1)（异常 %d）%s'
            % (len(align_bad), align_bad[:3] if align_bad else ''))
    if eqs:
        R.check(numbered > 0, '带编号公式 %d 条' % numbered)
        if formula_mode == 'omml':
            R.check(n_math == len(eqs),
                    '公式排成 Word 公式对象（%d/%d；回退的会在生成时给警告）'
                    % (n_math, len(eqs)))
            R.check(not font_bad,
                    '数学 run 用 %s（否则 ∑ ∫ 会渲染成方框）' % L.MATH_FONT)
            R.check(all(t.cell(0, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
                        for t in eqs), '公式栏居中')
        else:
            R.check(n_math == 0, '按 --formula text 出稿：不带公式对象（%d 个）' % n_math)


def check_parts(d, path, st, mode, R):
    print('\n=== 六部分齐全 ===')
    texts = [p.text.strip() for p in d.paragraphs]
    have = {
        SP.PART_COVER: bool(st['has_meta'] and st['meta_title']),
        SP.PART_ABSTRACT: any(t in ('摘要', '中文摘要', '摘  要') for t in texts),
        SP.PART_KEYWORDS: any(p.style.name == SP.STYLE_KW for p in d.paragraphs),
        SP.PART_TOC: b'TOC' in zipfile.ZipFile(path).read('word/document.xml'),
        SP.PART_BODY: st['n_chapters'] >= 1,
        SP.PART_REFS: any(t == '参考文献' for t in texts),
    }
    for part in SP.REQUIRED_PARTS:
        if mode == SP.MODE_DAILY and part in SP.OMIT_IN_DAILY:
            R.check(not have[part], '日常版已省略 %s' % part)
        elif mode == SP.MODE_DAILY and part == SP.PART_COVER:
            # 日常版不强制封面（与 md2paper.validate 的约定一致）：没写 front matter 就不出封面
            R.check(True, '封面：%s（日常版不强制）'
                    % ('有' if have[part] else '源稿无 front matter，不出封面'))
        else:
            R.check(have[part], '%s 齐全' % part)


def _parsed(md_path):
    """→ (blocks, meta, mode 占位, st) 供 plan_zones 使用。"""
    with open(md_path, encoding='utf-8') as f:
        text = f.read()
    meta, body = PP.split_front_matter(text)
    blocks = PP.retag(P.parse_blocks(body, is_path=False), bool(meta))
    st = PP.annotate(blocks, meta)
    return blocks, meta, None, st


def _docx_stream(d):
    """同 V.docx_stream，但公式段落要单独抠 oMath 里的 m:t。

    python-docx 的 `paragraph.text` 只遍历 w:r —— 数学区里的字它一个也看不见，
    照搬会把整条公式当成空的，内容无损校验直接假绿。
    """
    out = []
    for ch in d.element.body.iterchildren():
        if ch.tag == qn('w:p'):
            p = Paragraph(ch, d)
            if p.style.name == SP.STYLE_EQ:
                mt = ''.join(t.text or '' for t in ch.iter(qn('m:t')))
                if mt:
                    out.append((mt, False))
                    continue
            if p.text.strip():
                fi = p.paragraph_format.first_line_indent
                out.append((p.text, bool(fi is not None and fi < 0)))
        elif ch.tag == qn('w:tbl'):
            tbl = Table(ch, d)
            unit = _formula_table_unit(tbl)
            if unit is not None:
                out.append((unit, False))       # 公式表只产出公式那一个单元
                continue
            for r in tbl.rows:
                for c in r.cells:
                    for p in c.paragraphs:
                        if p.text.strip():
                            out.append((p.text, False))
    return out


def _formula_table_unit(tbl):
    """公式表 → 该产出的那一个文本单元（中栏的公式字符）；不是公式表返回 None。

    编号栏不能进文本流：它是生成器**加上去**的，源稿里没有（比对时按 EQNUM 剥掉）。
    """
    try:
        if len(tbl.rows) != 1 or len(tbl.columns) != 3:
            return None
        cell = tbl.cell(0, 1)
        if not any(p.style.name == SP.STYLE_EQ for p in cell.paragraphs):
            return None
        p = cell.paragraphs[0]
        return _math_text(p._p) or p.text or ''
    except Exception:
        return None


def check_content(d, md_path, mode, R, formula_mode='omml'):
    print('\n=== 内容无损 ===')
    nz = lambda s: re.sub(r'\s+', '', s or '')
    canon = lambda x: EQNUM.sub('', CAPNUM.sub(r'\1', x))

    blocks, meta, _, st = _parsed(md_path)
    A = [canon(nz(x)) for x in
         PP.stream_lines(blocks, meta, mode, st, formula_mode)]
    A = [x for x in A if x]
    B = []
    for x, is_list in _docx_stream(d):
        if V.TOC_HINT in x:
            continue
        x = nz(V.LIST_MARK.sub('', x) if is_list else x)
        if x:
            B.append(canon(x))
    R.check(A == B, '按 body 顺序的文本流逐单元一致（md %d vs docx %d）'
            % (len(A), len(B)))
    if A != B:
        i = next((i for i, (x, y) in enumerate(zip(A, B)) if x != y), min(len(A), len(B)))
        print('      #%d\n        md  =%r\n        docx=%r'
              % (i, (A[i] if i < len(A) else '<无>')[:88],
                 (B[i] if i < len(B) else '<无>')[:88]))
    sa, sb = nz(''.join(A)), nz(''.join(B))
    R.check(sa == sb, '全文去除格式后字符完全一致（%d vs %d）—— 内容零丢失'
            % (len(sa), len(sb)))
    if sa != sb:
        i = next((i for i, (x, y) in enumerate(zip(sa, sb)) if x != y), min(len(sa), len(sb)))
        print('      @%d md…%r\n         dc…%r'
              % (i, sa[max(0, i - 30):i + 30], sb[max(0, i - 30):i + 30]))


def main(argv=None):
    ap = argparse.ArgumentParser(description='校验 md2paper 生成的论文 docx')
    ap.add_argument('docx')
    ap.add_argument('--md', help='源 Markdown，给了就做内容无损校验')
    ap.add_argument('--mode', choices=list(SP.MODES), help='默认从 docx 里读')
    ap.add_argument('--formula', choices=['omml', 'text'], help='默认从 docx 里读')
    a = ap.parse_args(argv)

    d = Document(a.docx)
    tag = d.core_properties.comments or ''
    mode = a.mode or (_search(r'doc-paper-mode:(\w+)', tag))
    formula_mode = a.formula or _search(r'formula:(\w+)', tag) or 'omml'
    if mode not in SP.MODES:
        print('⚠ 读不出模式（docx 里没有 %s），也没有 --mode —— 按正式版校验'
              % (md2paper.CORE_MODE_TAG % 'paper'))
        mode = SP.MODE_PAPER

    R = V.Report()
    print('校验 %s（模式：%s，公式：%s）' % (a.docx, mode, formula_mode))
    check_styles(d, R)
    check_kw(d, R, mode)
    check_formula(d, R, formula_mode)
    if a.md:
        blocks, meta, _, st = _parsed(a.md)
        check_refs(d, st, R, mode)
        check_sections(d, a.docx, blocks, st, R, mode)
        check_cover(d, st, R)
        check_parts(d, a.docx, st, mode, R)
        # 按底层 XML 元素比 —— python-docx 每次访问 d.tables 都新建包装对象，
# 直接拿 Table 对象做 `in` 判断永远是 False（包装对象各自独立）
        ft_elems = {id(t._tbl) for t in formula_tables(d)}
        body_tables = [t for t in d.tables if id(t._tbl) not in ft_elems]
        if body_tables:
            # 公式表不是三线表，别拿去喂 doc-format 的表格断言
            V.check_tables(_TableView(d, body_tables), R)
        V.check_images(d, a.md, R)
        check_content(d, a.md, mode, R, formula_mode)
    print('\n段落 %d | 表格 %d | 节 %d' % (len(d.paragraphs), len(d.tables), len(d.sections)))
    return 1 if R.summary() else 0


if __name__ == '__main__':
    sys.exit(main())